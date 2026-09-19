"""
Bot Status Report workbook (the file emailed to users and dropped in SharePoint).

Same layout as the manual "Bot Status Report" files the team used to build:
  - "Bot Status"        pivot: count of runs by Status
  - "Bot Run List"      pivot: runs per bot (Automation name) by Occurrence
  - "Control Room Dump" the run records, with a Remarks column explaining
                        failures and status overrides

The pivot tables are real Excel pivots (openpyxl cannot create them, so the
pivot parts are written into the saved package). Their cells are pre-filled
with the same values Excel would show, so the report reads correctly in
Outlook preview / Protected View before any refresh.
"""

import os
import re
import io
import zipfile
from collections import Counter
from datetime import datetime, timedelta, timezone
from typing import Any, Callable, Dict, List, Optional
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

IST = timezone(timedelta(hours=5, minutes=30))

# Bots whose non-completed runs are reported as Completed (their failures are
# known/benign). Matched as a prefix of the automation name, case-insensitive.
FORCE_COMPLETED_BOTS = [
    p.strip().upper()
    for p in os.getenv("AA_FORCE_COMPLETED_BOTS", "AGEL014,AUC172").split(",")
    if p.strip()
]

COLUMNS = [
    "Status", "Automation type", "Activity name", "Running time",
    "Automation name", "Occurrence", "Device", "Run as user",
    "Activity type", "Started on", "Ended on", "Remarks",
]
# Remarks is kept out of the pivot source: pivot cache items are limited to
# 255 characters and error text can be longer.
PIVOT_SOURCE_COLUMNS = len(COLUMNS) - 1
COL = {name: i for i, name in enumerate(COLUMNS)}

SHEET_STATUS = "Bot Status"
SHEET_RUN_LIST = "Bot Run List"
SHEET_DUMP = "Control Room Dump"

COMPLETED = "Completed"

# Control Room API status -> label used in the Control Room UI exports
STATUS_LABELS = {
    "COMPLETED": COMPLETED,
    "RUN_FAILED": "Failed",
    "DEPLOY_FAILED": "Deploy failed",
    "RUN_ABORTED": "Stopped",
    "RUN_TIMED_OUT": "Timed out",
    "RUN_PAUSED": "Paused",
}
ACTIVITY_TYPE_LABELS = {
    "SCHEDULED": "Schedule",
    "TRIGGER": "Trigger",
    "RUN_NOW": "Run now",
}
AUTOMATION_TYPE_LABELS = {"BOT": "Task Bot"}
NOT_IN_MASTER = "Not in master"

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
FAILED_FILL = PatternFill(start_color="FDE2E1", end_color="FDE2E1", fill_type="solid")
OVERRIDE_FILL = PatternFill(start_color="FFF4CE", end_color="FFF4CE", fill_type="solid")


# ===========================================================================
# Record -> row
# ===========================================================================
def _label(value: Optional[str], mapping: Dict[str, str]) -> str:
    if not value:
        return ""
    return mapping.get(value, value.replace("_", " ").capitalize())


def _parse_utc(value: Optional[str]) -> Optional[datetime]:
    if not value:
        return None
    s = str(value).strip().replace("Z", "+00:00")
    s = re.sub(r"(\.\d{6})\d+", r"\1", s)  # AA sends 7 fractional digits
    try:
        dt = datetime.fromisoformat(s)
    except ValueError:
        return None
    return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)


def format_ist(dt: Optional[datetime]) -> str:
    """Format like the Control Room UI export: '2026-06-20 00:40:15 IST'."""
    return dt.astimezone(IST).strftime("%Y-%m-%d %H:%M:%S IST") if dt else ""


def _running_time(start: Optional[datetime], end: Optional[datetime]) -> str:
    if not start or not end or end < start:
        return ""
    total = int((end - start).total_seconds())
    hours, rem = divmod(total, 3600)
    minutes, seconds = divmod(rem, 60)
    if hours:
        return f"{hours}h {minutes}m {seconds}s"
    if minutes:
        return f"{minutes}m {seconds}s"
    return f"{seconds}s"


def _error_text(record: Dict[str, Any]) -> str:
    """One-line explanation from the AA error object."""
    err = record.get("error")
    if not isinstance(err, dict):
        return ""
    message = (err.get("message") or "").strip()
    detail = ""
    details_values = err.get("detailsValues") or []
    if details_values and isinstance(details_values[0], str):
        detail = details_values[0]
    elif err.get("details"):
        detail = err["details"].replace("This may be due to the following reason:", "")
    text = " - ".join(p for p in (message, detail.strip()) if p)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:500]


def _occurrence_label(schedule: Optional[str]) -> str:
    s = (schedule or "").strip().lower()
    if not s:
        return NOT_IN_MASTER
    if "demand" in s:
        return "On Demand"
    if "multiple" in s:
        return "Multiple Times a Day"
    if "daily" in s:
        return "Daily"
    if "week" in s:
        return "Weekly"
    if "month" in s:
        return "Monthly"
    return schedule.strip()


def _forced_completed_prefix(automation_name: str) -> Optional[str]:
    upper = automation_name.upper()
    for prefix in FORCE_COMPLETED_BOTS:
        if upper.startswith(prefix):
            return prefix
    return None


def build_rows(records: List[Dict[str, Any]],
               schedule_lookup: Callable[[str], Optional[str]]) -> List[List[str]]:
    """
    Convert AA activity records to report rows (in COLUMNS order), newest
    first, de-duplicated by execution id.

    schedule_lookup(automation_name) returns the bot's schedule from the
    master list, or None when the bot is not in the master.
    """
    seen = set()
    rows = []
    for rec in records:
        exec_id = rec.get("id") or (
            f"{rec.get('automationName')}_{rec.get('startDateTime')}"
            f"_{rec.get('endDateTime')}_{rec.get('deviceName')}"
        )
        if exec_id in seen:
            continue
        seen.add(exec_id)

        activity_name = rec.get("automationName") or ""
        automation_name = (rec.get("fileName") or rec.get("currentBotName")
                           or re.split(r"\.\d+", activity_name)[0] or "Unknown")
        start = _parse_utc(rec.get("startDateTime"))
        end = _parse_utc(rec.get("endDateTime"))

        api_status = rec.get("status") or ""
        status = _label(api_status, STATUS_LABELS) or "Unknown"
        error = _error_text(rec)
        remarks = ""
        if status != COMPLETED:
            prefix = _forced_completed_prefix(automation_name)
            if prefix:
                remarks = (f"Control Room status: {status}. Marked Completed as per "
                           f"CoBot rule for {prefix}.")
                if error:
                    remarks += f" Error: {error}"
                status = COMPLETED
            else:
                remarks = error or f"Control Room status: {status}"

        rows.append([
            status,
            _label(rec.get("activityType"), AUTOMATION_TYPE_LABELS),
            activity_name,
            _running_time(start, end),
            automation_name,
            _occurrence_label(schedule_lookup(automation_name) if automation_name else None),
            rec.get("deviceName") or "",
            rec.get("userName") or "",
            _label(rec.get("initiationType") or rec.get("type"), ACTIVITY_TYPE_LABELS),
            format_ist(start),
            format_ist(end),
            remarks,
        ])

    rows.sort(key=lambda r: r[COL["Ended on"]], reverse=True)
    return rows


# ===========================================================================
# Workbook
# ===========================================================================
def _write_title(ws, title: str, subtitle: str):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="595959")


def _render_status_pivot(ws, rows) -> Dict[str, Any]:
    counts = Counter(r[COL["Status"]] for r in rows)
    statuses = sorted(counts)
    ws.cell(3, 1, "Row Labels").font = Font(bold=True)
    ws.cell(3, 2, "Count of Status").font = Font(bold=True)
    for i, status in enumerate(statuses):
        ws.cell(4 + i, 1, status)
        ws.cell(4 + i, 2, counts[status])
    total_row = 4 + len(statuses)
    ws.cell(total_row, 1, "Grand Total").font = Font(bold=True)
    ws.cell(total_row, 2, len(rows)).font = Font(bold=True)
    ws.column_dimensions["A"].width = max(14, max((len(s) for s in statuses), default=0) + 4)
    ws.column_dimensions["B"].width = 16
    return {"ref": f"A3:B{total_row}", "row_items": statuses}


def _render_run_list_pivot(ws, rows) -> Dict[str, Any]:
    names = sorted({r[COL["Automation name"]] for r in rows})
    occurrences = sorted({r[COL["Occurrence"]] for r in rows})
    counts = Counter((r[COL["Automation name"]], r[COL["Occurrence"]]) for r in rows)
    bold = Font(bold=True)

    ws.cell(3, 1, "Count of Automation name").font = bold
    ws.cell(3, 2, "Column Labels").font = bold
    ws.cell(4, 1, "Row Labels").font = bold
    for j, occ in enumerate(occurrences):
        ws.cell(4, 2 + j, occ).font = bold
    total_col = 2 + len(occurrences)
    ws.cell(4, total_col, "Grand Total").font = bold

    for i, name in enumerate(names):
        r = 5 + i
        ws.cell(r, 1, name)
        row_total = 0
        for j, occ in enumerate(occurrences):
            n = counts.get((name, occ), 0)
            if n:
                ws.cell(r, 2 + j, n)
                row_total += n
        ws.cell(r, total_col, row_total)

    total_row = 5 + len(names)
    ws.cell(total_row, 1, "Grand Total").font = bold
    for j, occ in enumerate(occurrences):
        ws.cell(total_row, 2 + j, sum(counts.get((n, occ), 0) for n in names)).font = bold
    ws.cell(total_row, total_col, len(rows)).font = bold

    ws.column_dimensions["A"].width = min(80, max(26, max((len(n) for n in names), default=0) + 4))
    for j, occ in enumerate(occurrences):
        ws.column_dimensions[get_column_letter(2 + j)].width = max(12, len(occ) + 4)
    ws.column_dimensions[get_column_letter(total_col)].width = 13
    return {
        "ref": f"A3:{get_column_letter(total_col)}{total_row}",
        "row_items": names,
        "col_items": occurrences,
    }


def _write_dump(ws, rows):
    for c, header in enumerate(COLUMNS, 1):
        cell = ws.cell(1, c, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r, row in enumerate(rows, 2):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
        remarks = row[COL["Remarks"]]
        if remarks:
            ws.cell(r, COL["Remarks"] + 1).alignment = Alignment(wrap_text=True, vertical="top")
            fill = OVERRIDE_FILL if row[COL["Status"]] == COMPLETED else FAILED_FILL
            ws.cell(r, COL["Status"] + 1).fill = fill

    for c, header in enumerate(COLUMNS, 1):
        letter = get_column_letter(c)
        if header == "Remarks":
            ws.column_dimensions[letter].width = 80
            continue
        longest = max([len(header)] + [len(str(row[c - 1])) for row in rows])
        ws.column_dimensions[letter].width = min(60, max(10, longest + 2))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, len(rows)) + 1}"


def write_report(records: List[Dict[str, Any]], path: str, title: str,
                 window_text: str,
                 schedule_lookup: Callable[[str], Optional[str]]) -> Dict[str, Any]:
    """
    Build the Bot Status Report workbook at `path`.
    Returns {"rows": int, "status_counts": {...}, "overridden": int}.
    """
    rows = build_rows(records, schedule_lookup)
    generated = datetime.now(IST).strftime("%d %b %Y %H:%M IST")
    subtitle = f"Window: {window_text}  |  Generated: {generated}"

    wb = openpyxl.Workbook()
    ws_status = wb.active
    ws_status.title = SHEET_STATUS
    ws_runs = wb.create_sheet(SHEET_RUN_LIST)
    ws_dump = wb.create_sheet(SHEET_DUMP)

    for ws in (ws_status, ws_runs):
        _write_title(ws, title, subtitle)
    status_layout = _render_status_pivot(ws_status, rows)
    runs_layout = _render_run_list_pivot(ws_runs, rows)
    _write_dump(ws_dump, rows)

    buffer = io.BytesIO()
    wb.save(buffer)
    package = _add_pivot_tables(buffer.getvalue(), rows, status_layout, runs_layout)
    with open(path, "wb") as f:
        f.write(package)

    return {
        "rows": len(rows),
        "status_counts": dict(Counter(r[COL["Status"]] for r in rows)),
        "overridden": sum(1 for r in rows
                          if r[COL["Status"]] == COMPLETED and r[COL["Remarks"]]),
    }


# ===========================================================================
# Pivot table parts
# ===========================================================================
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
REL_PIVOT_TABLE = f"{NS_REL}/pivotTable"
REL_PIVOT_CACHE_DEF = f"{NS_REL}/pivotCacheDefinition"
REL_PIVOT_CACHE_RECORDS = f"{NS_REL}/pivotCacheRecords"
CT_BASE = "application/vnd.openxmlformats-officedocument.spreadsheetml"
CACHE_ID = 1

_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _attr(value: str) -> str:
    return escape(_ILLEGAL_XML.sub("", str(value)), {'"': "&quot;"})


def _pivot_cache_parts(rows, axis_fields: Dict[int, List[str]]):
    """Return (cacheDefinition xml, cacheRecords xml)."""
    record_count = len(rows)
    source_ref = f"A1:{get_column_letter(PIVOT_SOURCE_COLUMNS)}{record_count + 1}"
    index = {f: {v: i for i, v in enumerate(items)} for f, items in axis_fields.items()}

    fields_xml = []
    for f in range(PIVOT_SOURCE_COLUMNS):
        name = _attr(COLUMNS[f])
        if f in axis_fields:
            items = "".join(f'<s v="{_attr(v)}"/>' for v in axis_fields[f])
            shared = f'<sharedItems count="{len(axis_fields[f])}">{items}</sharedItems>'
        else:
            values = [r[f] for r in rows]
            has_blank = any(v == "" for v in values)
            has_text = any(v != "" for v in values)
            if has_text and has_blank:
                shared = '<sharedItems containsBlank="1"/>'
            elif has_blank:
                shared = '<sharedItems containsNonDate="0" containsString="0" containsBlank="1"/>'
            else:
                shared = "<sharedItems/>"
        fields_xml.append(f'<cacheField name="{name}" numFmtId="0">{shared}</cacheField>')

    definition = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<pivotCacheDefinition xmlns="{NS_MAIN}" xmlns:r="{NS_REL}" r:id="rId1" '
        'createdVersion="8" refreshedVersion="8" minRefreshableVersion="3" '
        f'recordCount="{record_count}">'
        f'<cacheSource type="worksheet"><worksheetSource ref="{source_ref}" '
        f'sheet="{_attr(SHEET_DUMP)}"/></cacheSource>'
        f'<cacheFields count="{PIVOT_SOURCE_COLUMNS}">{"".join(fields_xml)}</cacheFields>'
        '</pivotCacheDefinition>'
    )

    record_xml = []
    for r in rows:
        cells = []
        for f in range(PIVOT_SOURCE_COLUMNS):
            v = r[f]
            if f in index:
                cells.append(f'<x v="{index[f][v]}"/>')
            elif v == "":
                cells.append("<m/>")
            else:
                cells.append(f'<s v="{_attr(v)}"/>')
        record_xml.append(f'<r>{"".join(cells)}</r>')
    records = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<pivotCacheRecords xmlns="{NS_MAIN}" xmlns:r="{NS_REL}" count="{record_count}">'
        f'{"".join(record_xml)}</pivotCacheRecords>'
    )
    return definition, records


def _pivot_field_xml(field: int, axis_fields, row_field, col_field, data_field) -> str:
    attrs = ['showAll="0"']
    if field == data_field:
        attrs.insert(0, 'dataField="1"')
    if field == row_field:
        attrs.insert(0, 'axis="axisRow"')
    elif field == col_field:
        attrs.insert(0, 'axis="axisCol"')
    if field in (row_field, col_field):
        n = len(axis_fields[field])
        items = "".join(f'<item x="{i}"/>' for i in range(n))
        return (f'<pivotField {" ".join(attrs)} sortType="ascending">'
                f'<items count="{n + 1}">{items}<item t="default"/></items></pivotField>')
    return f'<pivotField {" ".join(attrs)}/>'


def _axis_items_xml(n: int) -> str:
    items = "".join("<i><x/></i>" if i == 0 else f'<i><x v="{i}"/></i>' for i in range(n))
    return f'{items}<i t="grand"><x/></i>'


def _pivot_table_xml(name: str, layout, axis_fields, row_field: int,
                     col_field: Optional[int], data_name: str) -> str:
    fields = "".join(
        _pivot_field_xml(f, axis_fields, row_field, col_field, row_field)
        for f in range(PIVOT_SOURCE_COLUMNS)
    )
    n_rows = len(axis_fields[row_field])
    row_part = (f'<rowFields count="1"><field x="{row_field}"/></rowFields>'
                f'<rowItems count="{n_rows + 1}">{_axis_items_xml(n_rows)}</rowItems>')
    if col_field is None:
        location = f'<location ref="{layout["ref"]}" firstHeaderRow="1" firstDataRow="1" firstDataCol="1"/>'
        col_part = '<colItems count="1"><i/></colItems>'
    else:
        n_cols = len(axis_fields[col_field])
        location = f'<location ref="{layout["ref"]}" firstHeaderRow="1" firstDataRow="2" firstDataCol="1"/>'
        col_part = (f'<colFields count="1"><field x="{col_field}"/></colFields>'
                    f'<colItems count="{n_cols + 1}">{_axis_items_xml(n_cols)}</colItems>')
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<pivotTableDefinition xmlns="{NS_MAIN}" name="{name}" cacheId="{CACHE_ID}" '
        'applyNumberFormats="0" applyBorderFormats="0" applyFontFormats="0" '
        'applyPatternFormats="0" applyAlignmentFormats="0" applyWidthHeightFormats="1" '
        'dataCaption="Values" updatedVersion="8" minRefreshableVersion="3" '
        'useAutoFormatting="1" itemPrintTitles="1" createdVersion="8" indent="0" '
        'outline="1" outlineData="1" multipleFieldFilters="0">'
        f'{location}'
        f'<pivotFields count="{PIVOT_SOURCE_COLUMNS}">{fields}</pivotFields>'
        f'{row_part}{col_part}'
        f'<dataFields count="1"><dataField name="{_attr(data_name)}" fld="{row_field}" '
        'subtotal="count" baseField="0" baseItem="0"/></dataFields>'
        '<pivotTableStyleInfo name="PivotStyleLight16" showRowHeaders="1" '
        'showColHeaders="1" showRowStripes="0" showColStripes="0" showLastColumn="1"/>'
        '</pivotTableDefinition>'
    )


def _next_rid(rels_xml: str) -> str:
    ids = [int(n) for n in re.findall(r'Id="rId(\d+)"', rels_xml)]
    return f"rId{max(ids, default=0) + 1}"


def _add_relationship(rels_xml: Optional[str], rel_type: str, target: str):
    if not rels_xml:
        rels_xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    f'<Relationships xmlns="{NS_PKG_REL}"></Relationships>')
    rid = _next_rid(rels_xml)
    rel = f'<Relationship Id="{rid}" Type="{rel_type}" Target="{target}"/>'
    return rels_xml.replace("</Relationships>", rel + "</Relationships>"), rid


def _sheet_part(files: Dict[str, bytes], sheet_name: str) -> str:
    """Resolve a sheet name to its part path, e.g. 'xl/worksheets/sheet1.xml'."""
    workbook = files["xl/workbook.xml"].decode("utf-8")
    rels = files["xl/_rels/workbook.xml.rels"].decode("utf-8")
    m = re.search(rf'<sheet [^>]*name="{re.escape(_attr(sheet_name))}"[^>]*r:id="(rId\d+)"', workbook)
    if not m:
        raise ValueError(f"Sheet not found in package: {sheet_name}")
    t = re.search(rf'<Relationship [^>]*Id="{m.group(1)}"[^>]*Target="([^"]+)"', rels) or \
        re.search(rf'<Relationship [^>]*Target="([^"]+)"[^>]*Id="{m.group(1)}"', rels)
    target = t.group(1).lstrip("/")
    return target if target.startswith("xl/") else f"xl/{target}"


def _add_pivot_tables(package: bytes, rows, status_layout, runs_layout) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(package))
    files = {n: zin.read(n) for n in zin.namelist()}
    order = zin.namelist()

    axis_fields = {
        COL["Status"]: status_layout["row_items"],
        COL["Automation name"]: runs_layout["row_items"],
        COL["Occurrence"]: runs_layout["col_items"],
    }
    cache_def, cache_records = _pivot_cache_parts(rows, axis_fields)
    new_parts = {
        "xl/pivotCache/pivotCacheDefinition1.xml": cache_def,
        "xl/pivotCache/pivotCacheRecords1.xml": cache_records,
        "xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Relationships xmlns="{NS_PKG_REL}"><Relationship Id="rId1" '
            f'Type="{REL_PIVOT_CACHE_RECORDS}" Target="pivotCacheRecords1.xml"/></Relationships>'
        ),
        "xl/pivotTables/pivotTable1.xml": _pivot_table_xml(
            "PivotTable1", status_layout, axis_fields, COL["Status"], None, "Count of Status"),
        "xl/pivotTables/pivotTable2.xml": _pivot_table_xml(
            "PivotTable2", runs_layout, axis_fields, COL["Automation name"],
            COL["Occurrence"], "Count of Automation name"),
    }
    for i in (1, 2):
        new_parts[f"xl/pivotTables/_rels/pivotTable{i}.xml.rels"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Relationships xmlns="{NS_PKG_REL}"><Relationship Id="rId1" '
            f'Type="{REL_PIVOT_CACHE_DEF}" Target="../pivotCache/pivotCacheDefinition1.xml"/>'
            '</Relationships>'
        )

    # Attach each pivot table to its sheet
    for i, sheet_name in ((1, SHEET_STATUS), (2, SHEET_RUN_LIST)):
        sheet_path = _sheet_part(files, sheet_name)
        folder, fname = sheet_path.rsplit("/", 1)
        rels_path = f"{folder}/_rels/{fname}.rels"
        existing = files.get(rels_path)
        rels_xml, _ = _add_relationship(existing.decode("utf-8") if existing else None,
                                        REL_PIVOT_TABLE, f"../pivotTables/pivotTable{i}.xml")
        files[rels_path] = rels_xml.encode("utf-8")
        if rels_path not in order:
            order.append(rels_path)

    # Register the cache with the workbook
    wb_rels, cache_rid = _add_relationship(
        files["xl/_rels/workbook.xml.rels"].decode("utf-8"),
        REL_PIVOT_CACHE_DEF, "pivotCache/pivotCacheDefinition1.xml")
    files["xl/_rels/workbook.xml.rels"] = wb_rels.encode("utf-8")

    workbook = files["xl/workbook.xml"].decode("utf-8")
    pivot_caches = f'<pivotCaches><pivotCache cacheId="{CACHE_ID}" r:id="{cache_rid}"/></pivotCaches>'
    # CT_Workbook order: ... definedNames, calcPr, oleSize, customWorkbookViews, pivotCaches ...
    m = re.search(r"<calcPr[^>]*/>|<calcPr[^>]*>.*?</calcPr>", workbook, re.S)
    if m:
        workbook = workbook[:m.end()] + pivot_caches + workbook[m.end():]
    else:
        workbook = workbook.replace("</workbook>", pivot_caches + "</workbook>")
    if 'xmlns:r="' not in workbook.split(">", 2)[1]:
        workbook = workbook.replace("<workbook ", f'<workbook xmlns:r="{NS_REL}" ', 1)
    files["xl/workbook.xml"] = workbook.encode("utf-8")

    content_types = files["[Content_Types].xml"].decode("utf-8")
    overrides = (
        f'<Override PartName="/xl/pivotCache/pivotCacheDefinition1.xml" '
        f'ContentType="{CT_BASE}.pivotCacheDefinition+xml"/>'
        f'<Override PartName="/xl/pivotCache/pivotCacheRecords1.xml" '
        f'ContentType="{CT_BASE}.pivotCacheRecords+xml"/>'
        f'<Override PartName="/xl/pivotTables/pivotTable1.xml" '
        f'ContentType="{CT_BASE}.pivotTable+xml"/>'
        f'<Override PartName="/xl/pivotTables/pivotTable2.xml" '
        f'ContentType="{CT_BASE}.pivotTable+xml"/>'
    )
    files["[Content_Types].xml"] = content_types.replace("</Types>", overrides + "</Types>").encode("utf-8")

    for path, xml in new_parts.items():
        files[path] = xml.encode("utf-8")
        order.append(path)

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in order:
            zout.writestr(name, files[name])
    return out.getvalue()


# ===========================================================================
# Reading a report back (for the email summary)
# ===========================================================================
def summarize_report(path: str) -> Dict[str, Any]:
    """
    Read a Bot Status Report workbook and return what the email needs:
    title, window, totals, status counts, overridden count and failed runs.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        title = subtitle = None
        if SHEET_STATUS in wb.sheetnames:
            first = list(wb[SHEET_STATUS].iter_rows(min_row=1, max_row=2, max_col=1, values_only=True))
            title = first[0][0] if first else None
            subtitle = first[1][0] if len(first) > 1 else None

        dump = wb[SHEET_DUMP] if SHEET_DUMP in wb.sheetnames else wb.worksheets[-1]
        rows = list(dump.iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        return {"title": title, "subtitle": subtitle, "total_runs": 0, "status_counts": {},
                "overridden": 0, "failed_runs": [], "bots": 0}

    header = {str(h).strip().lower(): i for i, h in enumerate(rows[0]) if h}

    def get(row, name):
        i = header.get(name.lower())
        return row[i] if i is not None and i < len(row) and row[i] is not None else ""

    def status_of(row):
        # Older AA exports hold raw API statuses (COMPLETED, RUN_FAILED...)
        raw = str(get(row, "Status") or "Unknown")
        return STATUS_LABELS.get(raw, raw)

    data = [r for r in rows[1:] if any(v is not None for v in r)]
    status_counts = Counter(status_of(r) for r in data)
    failed = [
        {
            "automation_name": get(r, "Automation name") or get(r, "Activity name"),
            "status": status_of(r),
            "ended_on": get(r, "Ended on"),
            "device": get(r, "Device") or get(r, "Device name"),
            "remarks": get(r, "Remarks") or get(r, "Error message"),
        }
        for r in data if status_of(r) != COMPLETED
    ]
    return {
        "title": title or f"Bot Status Report - {os.path.splitext(os.path.basename(path))[0]}",
        "subtitle": subtitle,
        "total_runs": len(data),
        "status_counts": dict(status_counts),
        "overridden": sum(1 for r in data
                          if status_of(r) == COMPLETED and get(r, "Remarks")),
        "failed_runs": failed,
        "bots": len({get(r, "Automation name") or get(r, "Activity name") for r in data}),
    }
