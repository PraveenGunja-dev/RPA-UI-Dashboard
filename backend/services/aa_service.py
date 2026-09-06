"""
Automation Anywhere Control Room API Service.

Connects to the AA Control Room, retrieves bot execution records for a
configurable time window, filters to target devices, and exports an Excel
file in the exact format the dashboard's SharePoint sync expects.

Usage (standalone):
    python -m services.aa_service            # last 12 hours
    python -m services.aa_service --hours 24 # last 24 hours

When imported, call  aa_sync(hours=12)  to run the full pipeline.
"""

import os
import json
import logging
import argparse
from collections import Counter
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple

import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from tenacity import retry, stop_after_attempt, wait_exponential, retry_if_exception_type
from requests.exceptions import RequestException
from dotenv import load_dotenv

# ---------------------------------------------------------------------------
# Environment
# ---------------------------------------------------------------------------
_env_path = Path(__file__).resolve().parent.parent / ".env"
load_dotenv(dotenv_path=_env_path)

# ---------------------------------------------------------------------------
# Logger  (rotating file + console)
# ---------------------------------------------------------------------------
from logging.handlers import RotatingFileHandler

_log_dir = Path(__file__).resolve().parent.parent / "logs" / "aa_sync"
_log_dir.mkdir(parents=True, exist_ok=True)

logger = logging.getLogger("AASync")
logger.setLevel(logging.INFO)

_fmt = logging.Formatter("%(asctime)s - %(levelname)s - %(message)s")

if not logger.handlers:
    _fh = RotatingFileHandler(
        str(_log_dir / "aa_sync.log"), maxBytes=5 * 1024 * 1024, backupCount=5
    )
    _fh.setFormatter(_fmt)
    logger.addHandler(_fh)

    _ch = logging.StreamHandler()
    _ch.setFormatter(_fmt)
    logger.addHandler(_ch)

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
AA_BASE_URL = os.getenv(
    "AA_BASE_URL", "https://adani-crprod.my.automationanywhere.digital"
).rstrip("/")
AA_USERNAME = os.getenv("AA_USERNAME", "")
AA_PASSWORD = os.getenv("AA_PASSWORD", "")
AA_API_KEY = os.getenv("AA_API_KEY", "")

# Target devices (the 7 runner machines used by Cobot bots)
TARGET_DEVICES = [
    "DT-1001-7B3NM24",
    "DT-1001-G62J7W3",
    "gcp10wbtrpap01",
    "gcp10wbtrpap02",
    "gcp10wbtrpap03",
    "gcp10wbtrpap04",
    "gcp10wbtrpap05",
]

# Output directory for exports
OUTPUT_DIR = Path(__file__).resolve().parent.parent / "uploads" / "aa_sync"


# ===========================================================================
# Authentication
# ===========================================================================
class AuthenticationError(Exception):
    """Raised when AA authentication fails for non-retryable reasons."""


@retry(
    stop=stop_after_attempt(3),
    wait=wait_exponential(multiplier=1, min=2, max=10),
    retry=retry_if_exception_type(RequestException),
)
def _authenticate() -> str:
    """
    Authenticates with Automation Anywhere Control Room.
    Returns the bearer token string.
    """
    if not AA_USERNAME or not AA_PASSWORD or not AA_API_KEY:
        raise AuthenticationError(
            "Missing AA credentials. Set AA_USERNAME, AA_PASSWORD, AA_API_KEY in .env"
        )

    url = f"{AA_BASE_URL}/v2/authentication"
    payload = {
        "username": AA_USERNAME,
        "password": AA_PASSWORD,
        "apikey": AA_API_KEY,
        "multipleLogin": False,
    }

    logger.info("Authenticating with Automation Anywhere Control Room...")
    resp = requests.post(url, json=payload, timeout=30)

    if resp.status_code == 200:
        token = resp.json().get("token")
        if token:
            logger.info("Authentication successful.")
            return token
        raise AuthenticationError("No token in authentication response.")
    elif resp.status_code in (401, 403):
        raise AuthenticationError(
            f"HTTP {resp.status_code} — check AA credentials or permissions."
        )
    else:
        resp.raise_for_status()  # let tenacity retry on 5xx


# ===========================================================================
# Activity Retrieval (with pagination + device filtering)
# ===========================================================================
@retry(
    stop=stop_after_attempt(5),
    wait=wait_exponential(multiplier=1, min=2, max=10),
    retry=retry_if_exception_type(RequestException),
)
def _fetch_page(token: str, offset: int, length: int,
                start_time: str, end_time: str) -> Dict[str, Any]:
    """Fetch a single page of activity records from the AA API."""
    url = f"{AA_BASE_URL}/v3/activity/list"
    headers = {"X-Authorization": token, "Content-Type": "application/json"}
    payload = {
        "filter": {
            "operator": "and",
            "filters": [
                {
                    "operator": "between",
                    "field": "endDateTime",
                    "value": [start_time, end_time],
                }
            ],
        },
        "sort": [{"field": "endDateTime", "direction": "desc"}],
        "page": {"length": length, "offset": offset},
    }

    resp = requests.post(url, headers=headers, json=payload, timeout=60)

    # Handle token expiry mid-pagination
    if resp.status_code == 401:
        logger.warning("HTTP 401 — re-authenticating...")
        new_token = _authenticate()
        headers["X-Authorization"] = new_token
        resp = requests.post(url, headers=headers, json=payload, timeout=60)

    resp.raise_for_status()
    return resp.json()


def _get_all_activities(token: str, start_time: str,
                        end_time: str) -> List[Dict[str, Any]]:
    """
    Paginates through all AA activities in the time window and filters
    locally for the target devices.
    """
    target_set = set(TARGET_DEVICES)
    filtered: List[Dict[str, Any]] = []
    offset = 0
    page_size = 200  # max safe page size

    logger.info(f"Fetching activities from {start_time} to {end_time}")

    total_expected = None

    while True:
        try:
            data = _fetch_page(token, offset, page_size, start_time, end_time)
        except Exception as e:
            logger.error(
                f"Failed to fetch activities at offset {offset}: {e}"
            )
            break

        items = data.get("list", [])

        if total_expected is None:
            total_expected = data.get("page", {}).get("totalFilter", 0)
            logger.info(f"Total API records in time window: {total_expected}")

        if not items:
            break

        out_of_bounds = False
        for item in items:
            end_dt = item.get("endDateTime", "")
            if end_dt and end_dt < start_time:
                out_of_bounds = True
                continue
            if item.get("deviceName") in target_set and end_dt <= end_time:
                filtered.append(item)

        if out_of_bounds:
            logger.info("Reached records older than start_time. Stopping.")
            break

        logger.info(f"Processed offset {offset}. Filtered so far: {len(filtered)}")

        if len(items) < page_size:
            break

        offset += page_size

    logger.info(
        f"Retrieval complete. {len(filtered)} records for "
        f"{len(target_set)} target devices."
    )
    return filtered


# ===========================================================================
# Excel Export  (SharePoint-compatible format)
# ===========================================================================
def _export_excel(records: List[Dict[str, Any]], start_time: str,
                  end_time: str) -> Optional[str]:
    """
    Generates a formatted .xlsx file with the exact column headers the
    backend's daily-report parser (excel_parser.py) expects, so it can be
    dropped straight into SharePoint.

    Returns the absolute path to the generated Excel file, or None on failure.
    """
    if not records:
        logger.warning("No records to export.")
        return None

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d_%H-%M-%S")
    excel_path = str(OUTPUT_DIR / f"Adani_Daily_Bot_Status_{timestamp}.xlsx")

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Control Room Dump"

        # Column headers that match what excel_parser.parse_daily_report expects
        headers = [
            "Activity Name",
            "Status",
            "Started On",
            "Ended On",
            "Device Name",
            "Automation Type",
            "Error Message",
        ]

        # Style headers
        hdr_font = Font(bold=True, color="FFFFFF")
        hdr_fill = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid"
        )
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[
                openpyxl.utils.get_column_letter(col_idx)
            ].width = 30

        # De-duplicate by execution ID and write rows
        seen_ids: set = set()
        row_idx = 2
        for rec in records:
            exec_id = rec.get("id") or (
                f"{rec.get('automationName', '')}_{rec.get('startDateTime', '')}"
                f"_{rec.get('endDateTime', '')}_{rec.get('deviceName', '')}"
            )
            if exec_id in seen_ids:
                continue
            seen_ids.add(exec_id)

            err = rec.get("error", {})
            error_msg = err.get("message", "") if isinstance(err, dict) else ""

            row_data = [
                rec.get("automationName", ""),
                rec.get("status", ""),
                rec.get("startDateTime", ""),
                rec.get("endDateTime", ""),
                rec.get("deviceName", ""),
                rec.get("automationType", ""),
                error_msg,
            ]
            for col_idx, val in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=val)
            row_idx += 1

        wb.save(excel_path)
        logger.info(f"Excel saved: {excel_path}  ({row_idx - 2} rows)")
        return excel_path

    except Exception as e:
        logger.error(f"Failed to write Excel: {e}")
        return None


# ===========================================================================
# JSON Export  (raw archive for debugging)
# ===========================================================================
def _export_json(records: List[Dict[str, Any]], start_time: str,
                 end_time: str) -> Optional[str]:
    """Saves the raw API records as a JSON archive."""
    if not records:
        return None

    json_dir = OUTPUT_DIR / "json"
    json_dir.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d_%H-%M-%S")
    json_path = str(json_dir / f"aa_raw_{timestamp}.json")

    try:
        payload = {
            "query_start_time": start_time,
            "query_end_time": end_time,
            "retrieval_timestamp": datetime.now(timezone.utc).isoformat(),
            "total_records": len(records),
            "records": records,
        }
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(payload, f, indent=2)
        logger.info(f"JSON archive saved: {json_path}")
        return json_path
    except Exception as e:
        logger.error(f"Failed to write JSON: {e}")
        return None


# ===========================================================================
# Public API — the single entry-point for the sync pipeline
# ===========================================================================
def aa_sync(
    hours: int = 12,
    start: str = None,
    end: str = None,
    upload_to_sharepoint: bool = True,
) -> Dict[str, Any]:
    """
    Runs the full Automation Anywhere sync pipeline:
      1. Authenticate with AA Control Room
      2. Fetch activity records (paginated, device-filtered)
      3. Export Excel (SharePoint format) + JSON archive
      4. Optionally upload the Excel file to SharePoint

    Returns a summary dict with status and file paths.
    """
    result: Dict[str, Any] = {
        "status": "error",
        "records": 0,
        "excel_path": None,
        "json_path": None,
        "sharepoint_upload": None,
    }

    # --- Resolve time window ---
    if start and end:
        start_time, end_time = start, end
    else:
        now = datetime.now(timezone.utc)
        start_time = (now - timedelta(hours=hours)).strftime("%Y-%m-%dT%H:%M:%SZ")
        end_time = now.strftime("%Y-%m-%dT%H:%M:%SZ")

    logger.info("=" * 60)
    logger.info("Automation Anywhere Sync -- Starting")
    logger.info(f"Time window: {start_time} -> {end_time}")
    logger.info("=" * 60)

    # --- 1. Authenticate ---
    try:
        token = _authenticate()
    except AuthenticationError as e:
        logger.error(f"Authentication failed: {e}")
        result["message"] = str(e)
        return result

    # --- 2. Fetch activities ---
    records = _get_all_activities(token, start_time, end_time)
    result["records"] = len(records)

    if not records:
        logger.warning("No execution records found for target devices.")
        result["status"] = "success"
        result["message"] = "No records found in time window."
        return result

    # Log summary
    status_counts = Counter(r.get("status", "UNKNOWN") for r in records)
    device_counts = Counter(r.get("deviceName", "UNKNOWN") for r in records)
    logger.info(f"Status breakdown: {dict(status_counts)}")
    logger.info(f"Device breakdown: {dict(device_counts)}")

    # --- 3. Export ---
    result["excel_path"] = _export_excel(records, start_time, end_time)
    result["json_path"] = _export_json(records, start_time, end_time)

    # --- 4. Upload to SharePoint ---
    if upload_to_sharepoint and result["excel_path"]:
        try:
            from services.sharepoint_service import SharePointService

            sp = SharePointService()
            upload_ok = sp.upload_file(result["excel_path"])
            result["sharepoint_upload"] = "success" if upload_ok else "failed"
            if upload_ok:
                logger.info("SharePoint upload: SUCCESS")
            else:
                logger.warning("SharePoint upload: FAILED (see logs above)")
        except Exception as e:
            result["sharepoint_upload"] = f"error: {e}"
            logger.error(f"SharePoint upload error: {e}")
    else:
        result["sharepoint_upload"] = "skipped"

    result["status"] = "success"
    logger.info("=" * 60)
    logger.info("Automation Anywhere Sync -- Complete")
    logger.info(f"Records: {result['records']}  |  Excel: {result['excel_path']}")
    logger.info(f"SharePoint: {result['sharepoint_upload']}")
    logger.info("=" * 60)

    return result


# ===========================================================================
# CLI entry-point  (python -m services.aa_service --hours 12)
# ===========================================================================
def _cli():
    parser = argparse.ArgumentParser(
        description="Automation Anywhere -> SharePoint Sync"
    )
    parser.add_argument(
        "--hours", type=int, default=12,
        help="Rolling hours window (default: 12)"
    )
    parser.add_argument("--start", type=str, help="Start ISO-8601 datetime")
    parser.add_argument("--end", type=str, help="End ISO-8601 datetime")
    parser.add_argument(
        "--no-upload", action="store_true",
        help="Skip SharePoint upload (export only)"
    )
    args = parser.parse_args()

    result = aa_sync(
        hours=args.hours,
        start=args.start,
        end=args.end,
        upload_to_sharepoint=not args.no_upload,
    )

    # Print human-readable summary
    print("\n" + "=" * 55)
    print("  Automation Anywhere -> SharePoint Sync  Summary")
    print("=" * 55)
    print(f"  Status:            {result['status']}")
    print(f"  Records fetched:   {result['records']}")
    print(f"  Excel file:        {result.get('excel_path', 'N/A')}")
    print(f"  JSON archive:      {result.get('json_path', 'N/A')}")
    print(f"  SharePoint upload: {result.get('sharepoint_upload', 'N/A')}")
    print("=" * 55 + "\n")


if __name__ == "__main__":
    _cli()
